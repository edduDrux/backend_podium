import re
from pathlib import Path

from pypdf import PdfReader


def _extract_text_from_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    parts: list[str] = []

    for page in reader.pages:
        txt = (page.extract_text() or "").strip()
        if txt:
            parts.append(txt)

    return "\n\n".join(parts)


def _extract_text_from_pptx(file_path: Path) -> str:
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(file_path))
    parts: list[str] = []

    for slide in prs.slides:
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        slide_texts.append(line)
        if slide_texts:
            parts.append("\n".join(slide_texts))

    return "\n\n".join(parts)


def _extract_text_from_docx(file_path: Path) -> str:
    from docx import Document as DocxDocument  # type: ignore

    doc = DocxDocument(str(file_path))
    parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

_SUFFIX_MAP: dict[str, callable] = {
    ".pdf": _extract_text_from_pdf,
    ".pptx": _extract_text_from_pptx,
    ".ppt": _extract_text_from_pptx,
    ".docx": _extract_text_from_docx,
    ".doc": _extract_text_from_docx,
}

_CONTENT_TYPE_MAP: dict[str, callable] = {
    "application/pdf": _extract_text_from_pdf,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": _extract_text_from_pptx,
    "application/vnd.ms-powerpoint": _extract_text_from_pptx,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": _extract_text_from_docx,
    "application/msword": _extract_text_from_docx,
}


def extract_text(path: str, content_type: str) -> str:
    """Detecta tipo pelo sufixo do arquivo, com fallback para content_type."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    extractor = _SUFFIX_MAP.get(suffix) or _CONTENT_TYPE_MAP.get(content_type)
    if extractor is None:
        raise ValueError(f"Tipo de arquivo não suportado: {suffix} ({content_type})")

    return extractor(file_path)


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 200) -> list[str]:
    """
    MVP: chunking por caracteres com overlap.
    Depois, se quiser, dá para migrar para chunking por tokens.
    """
    cleaned = re.sub(r"\s+", " ", text).strip()
    if not cleaned:
        return []

    chunks: list[str] = []
    start = 0
    n = len(cleaned)

    while start < n:
        end = min(start + max_chars, n)

        # tenta cortar em um ponto final próximo do final do chunk
        cut = cleaned.rfind(".", start, end)
        if cut == -1 or cut < start + int(max_chars * 0.5):
            cut = end
        else:
            cut = cut + 1  # inclui o ponto

        chunk = cleaned[start:cut].strip()
        if chunk:
            chunks.append(chunk)

        next_start = cut - overlap
        if next_start <= start:
            next_start = cut  # garante progresso
        start = next_start

    return chunks
